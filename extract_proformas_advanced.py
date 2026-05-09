# extract_proformas_advanced.py

import os
import re
import json
import time
import argparse
from pathlib import Path
from typing import Optional, Dict, List, Tuple

import docx
import fitz
import pytesseract
from PIL import Image
import pandas as pd

# ---------------- CONFIG ----------------

DATA_FOLDER = "datos_proformas"
OUT_FOLDER = "."

# Tesseract Windows
if os.name == "nt":
    default_tess = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    if os.path.exists(default_tess):
        pytesseract.pytesseract.tesseract_cmd = default_tess


# ---------------- UTILIDADES ----------------

def now_ts() -> str:
    return time.strftime("%Y%m%d_%H%M%S")


def clean_number(txt: str) -> Optional[float]:

    if txt is None:
        return None

    txt = str(txt)

    txt = (
        txt.replace("S/.", "")
        .replace("S/", "")
        .replace("s/", "")
        .replace("$", "")
    )

    txt = txt.strip()

    if txt.count(".") > 1 and txt.count(",") == 1:
        txt = txt.replace(".", "").replace(",", ".")

    elif txt.count(",") > 1 and txt.count(".") == 1:
        txt = txt.replace(",", "").replace(".", ".")

    else:
        if txt.count(",") == 1 and txt.count(".") == 0:
            txt = txt.replace(",", ".")
        else:
            txt = txt.replace(",", "")

    txt = re.sub(r"[^\d\.-]", "", txt)

    try:
        return float(txt)
    except:
        return None


# ---------------- LECTURA DOCX ----------------

def read_docx_text(path: str) -> str:

    doc = docx.Document(path)

    parts = []

    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())

    for table in doc.tables:
        for row in table.rows:

            row_text = " | ".join(
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            )

            if row_text:
                parts.append(row_text)

    return "\n".join(parts)


# ---------------- LECTURA PDF ----------------

def read_pdf_text(path: str) -> Tuple[str, bool]:

    doc = fitz.open(path)

    fulltext = ""

    page_images_needed = []

    for i, page in enumerate(doc):

        txt = page.get_text("text")

        if txt.strip():
            fulltext += txt + "\n"
        else:
            page_images_needed.append(i)

    if page_images_needed:

        ocr_text = []

        for i in page_images_needed:

            page = doc[i]

            pix = page.get_pixmap(dpi=300)

            mode = "RGBA" if pix.alpha else "RGB"

            img = Image.frombytes(
                mode,
                [pix.width, pix.height],
                pix.samples
            )

            txt = pytesseract.image_to_string(
                img,
                lang="spa+eng"
            )

            ocr_text.append(txt)

        fulltext += "\n".join(ocr_text)

        return fulltext, True

    return fulltext, False


# ---------------- REGEX ----------------

NUMBER_RE = r"([0-9]{1,3}(?:[.,][0-9]{3})*(?:[.,][0-9]{1,2})|[0-9]+(?:[.,][0-9]+)?)"

COST_PATTERNS = [
    rf"(?:costo\s*total|total\s*general|importe\s*total|total\s*a\s*pagar|total)[:\s\-]*S?[/.]?\s*{NUMBER_RE}",
]

MATERIALS_PATTERNS = [
    rf"(?:costo\s*de\s*materiales|materiales|material)[:\s\-]*S?[/.]?\s*{NUMBER_RE}"
]

TIME_PATTERNS = [
    rf"(?:tiempo\s*de\s*entrega|plazo|duraci[oó]n)[:\s\-]*{NUMBER_RE}",
    rf"(\d+)\s*(?:d[ií]as|dias)"
]


def search_patterns(text: str, patterns: List[str]) -> Optional[float]:

    for p in patterns:

        m = re.search(p, text, re.IGNORECASE)

        if m:

            groups = [g for g in m.groups() if g]

            if groups:

                val = clean_number(groups[-1])

                if val is not None:
                    return val

    return None


# ---------------- EXTRAER ----------------

def extract_from_text(text: str) -> Dict:

    reasons = []

    text = re.sub(r"\s+", " ", text)

    costo_total = search_patterns(text, COST_PATTERNS)

    costo_mat = search_patterns(text, MATERIALS_PATTERNS)

    tiempo = search_patterns(text, TIME_PATTERNS)

    if costo_total is None:

        reasons.append("sin_costo_total")

        return {
            "costoTotal": None,
            "costoMateriales": costo_mat,
            "tiempoEntrega": tiempo,
            "roi": None,
            "mbb": None,
            "reasons": reasons
        }

    if costo_mat and costo_mat > costo_total:
        costo_mat = None

    # ---------------- ROI Y MBB REALES ----------------

    roi = None
    mbb = None

    if costo_mat and costo_mat > 0:

        ganancia = costo_total - costo_mat

        roi = round(
            (ganancia / costo_mat) * 100,
            2
        )

        mbb = round(
            (ganancia / costo_total) * 100,
            2
        )

    return {

        "costoTotal": round(costo_total, 2),

        "costoMateriales": (
            round(costo_mat, 2)
            if costo_mat else None
        ),

        "tiempoEntrega": (
            round(tiempo, 2)
            if tiempo else None
        ),

        "roi": roi,

        "mbb": mbb,

        "reasons": reasons
    }


# ---------------- TABLAS DOCX ----------------

def try_extract_from_tables_docx(path: str) -> Optional[float]:

    try:

        doc = docx.Document(path)

        for table in doc.tables:

            headers = [
                cell.text.strip().lower()
                for cell in table.rows[0].cells
            ]

            total_col_idx = None

            for idx, h in enumerate(headers):

                if (
                    "total" in h
                    or "importe" in h
                    or "subtotal" in h
                ):
                    total_col_idx = idx
                    break

            if total_col_idx is not None:

                s = 0.0

                any_val = False

                for r in table.rows[1:]:

                    txt = r.cells[total_col_idx].text.strip()

                    val = clean_number(txt)

                    if val is not None:
                        s += val
                        any_val = True

                if any_val:
                    return s

    except:
        return None

    return None


# ---------------- MAIN ----------------

def run():

    base = Path(".").resolve()

    data_dir = base / DATA_FOLDER

    if not data_dir.exists():
        print("No existe carpeta datos_proformas")
        return

    ts = now_ts()

    out_csv = Path(OUT_FOLDER) / f"proformas_dataset_{ts}.csv"

    registros = []

    files = sorted([
        f for f in os.listdir(data_dir)
        if f.lower().endswith((".docx", ".pdf"))
    ])

    if not files:
        print("No hay archivos")
        return

    for f in files:

        path = data_dir / f

        print(f"\nProcesando: {f}")

        text = ""

        try:

            if f.lower().endswith(".docx"):

                table_sum = try_extract_from_tables_docx(str(path))

                text = read_docx_text(str(path))

                data = extract_from_text(text)

                if data["costoTotal"] is None and table_sum:

                    data = extract_from_text(
                        f"total: {table_sum}\n{text}"
                    )

            else:

                text, _ = read_pdf_text(str(path))

                data = extract_from_text(text)

        except Exception as e:

            print("ERROR:", e)

            continue

        descripcion = (
            text[:400].replace("\n", " ")
            if text else ""
        )

        registros.append({

            "archivo": f,

            "descripcion": descripcion,

            "costoTotal": data.get("costoTotal"),

            "costoMateriales": data.get("costoMateriales"),

            "tiempoEntrega": data.get("tiempoEntrega"),

            "roi": data.get("roi"),

            "mbb": data.get("mbb"),

            "factible": None,

            "reasons": ";".join(
                data.get("reasons", [])
            )
        })

        debug = {

            "archivo": f,

            "costoTotal": data.get("costoTotal"),

            "costoMateriales": data.get("costoMateriales"),

            "tiempoEntrega": data.get("tiempoEntrega"),

            "roi": data.get("roi"),

            "mbb": data.get("mbb")
        }

        debug_path = base / f"debug_{Path(f).stem}.json"

        with open(
            debug_path,
            "w",
            encoding="utf-8"
        ) as fh:

            json.dump(
                debug,
                fh,
                ensure_ascii=False,
                indent=2
            )

    df = pd.DataFrame(registros)

    df.to_csv(out_csv, index=False)

    print("\nColumnas generadas:")
    print(list(df.columns))

    print(f"\nDataset guardado en:")
    print(out_csv)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    args = parser.parse_args()
    run()
